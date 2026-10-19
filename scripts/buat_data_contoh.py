"""
Membuat ulang file contoh untuk demo:
- data/contoh_data_pekerja.xlsx   (data pekerja FIKTIF)
- templates/contoh_template_pkwt.docx (template kontrak GENERIK)

Semua nama, NIK, nomor pekerja, perusahaan, dan proyek di file ini
adalah karangan. NIK sengaja diawali "9999" (bukan kode wilayah yang ada).

Cara pakai (dari folder utama proyek):
    python scripts/buat_data_contoh.py
"""

from __future__ import annotations

import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

from config.settings import SAMPLE_EXCEL, SAMPLE_TEMPLATE  # noqa: E402
from scripts.template_builder import build_sample_template  # noqa: E402
from utils.formatters import format_indonesian_date  # noqa: E402

random.seed(2026)  # hasil selalu sama setiap kali dijalankan

COMPANY_CODE = "CKN"
MANAGER = "Rina Kartikasari"
NAMES = [
    "Andi Pratama", "Siti Rahmawati", "Budi Santoso", "Dewi Lestari", "Eko Saputra",
    "Fitri Handayani", "Gilang Ramadhan", "Hana Puspita", "Irfan Maulana", "Joko Susilo",
    "Kartika Sari", "Lukman Hakim", "Maya Anggraini", "Nanda Prakoso", "Oktaviani Putri",
    "Putra Wijaya", "Qori Amalia", "Rizky Firmansyah", "Sari Wulandari", "Taufik Hidayat",
    "Umar Fauzi", "Vina Melati", "Wahyu Nugroho", "Yuni Astuti", "Zaki Mubarak",
]
PROJECTS = [
    ("Proyek Pemeliharaan Tangki Timbun Area Contoh", "IO26001", "Cilacap"),
    ("Proyek Perawatan Jalur Pipa Wilayah Barat", "IO26002", "Balikpapan"),
    ("Proyek Overhaul Unit Kompresor Contoh", "IO26003", "Dumai"),
]
POSITIONS = ["Welder", "Rigger", "Safety Officer", "Admin Proyek", "Mechanic", "Electrician", "QC Inspector"]
BIRTH_CITIES = ["Jakarta", "Bandung", "Semarang", "Surabaya", "Medan", "Makassar", "Palembang", "Depok"]
STREETS = ["Jl. Melati", "Jl. Kenanga", "Jl. Cempaka", "Jl. Anggrek", "Jl. Flamboyan"]

HEADERS = [
    "NO", "NO KONTRAK", "WAKTU PEMBUATAN KONTRAK", "MANAGER", "NAMA PEKERJA", "TTL PEKERJA",
    "ALAMAT PEKERJA", "NIK PEKERJA", "NAMA PROYEK", "NO IO PROYEK", "NOPEK", "JABATAN PEKERJA",
    "TMT AWAL", "TMT AKHIR", "UPAH TETAP", "TUNJANGAN DAERAH", "TUNJANGAN LAINNYA",
    "LOKASI KERJA", "NAMA BAWAH",
]


def _valid_row(number: int, name: str) -> dict[str, object]:
    """Satu baris data pekerja fiktif yang lolos validasi."""
    project, io_number, location = PROJECTS[(number - 1) % len(PROJECTS)]
    start = date(2026, 1 + (number - 1) % 6, 1)
    end = start + timedelta(days=random.choice([180, 270, 364]))
    created = start - timedelta(days=random.randint(3, 10))
    birth = date(random.randint(1985, 2001), random.randint(1, 12), random.randint(1, 28))
    return {
        "NO": number,
        "NO KONTRAK": f"{number:03d}/PKWT/{COMPANY_CODE}/{start.month:02d}/2026",
        "WAKTU PEMBUATAN KONTRAK": format_indonesian_date(created, with_weekday=True),
        "MANAGER": MANAGER,
        "NAMA PEKERJA": name.upper(),
        "TTL PEKERJA": f"{random.choice(BIRTH_CITIES)}, {format_indonesian_date(birth)}",
        "ALAMAT PEKERJA": f"{random.choice(STREETS)} No. {random.randint(1, 99)}, "
                          f"RT 00{random.randint(1, 9)}/RW 00{random.randint(1, 9)}, Kota Contoh",
        "NIK PEKERJA": "9999" + "".join(random.choices("0123456789", k=12)),
        "NAMA PROYEK": project,
        "NO IO PROYEK": io_number,
        "NOPEK": f"9{number:07d}",
        "JABATAN PEKERJA": random.choice(POSITIONS),
        "TMT AWAL": format_indonesian_date(start),
        "TMT AKHIR": format_indonesian_date(end),
        "UPAH TETAP": random.choice([5_200_000, 5_750_000, 6_100_000, 6_800_000, 7_500_000]),
        "TUNJANGAN DAERAH": random.choice([0, 750_000, 1_000_000]),
        "TUNJANGAN LAINNYA": random.choice([None, 500_000]),
        "LOKASI KERJA": location,
        "NAMA BAWAH": name.upper(),
    }


def _sample_rows() -> list[dict[str, object]]:
    """20 data benar + 5 data yang SENGAJA salah untuk demo fitur validasi."""
    rows = [_valid_row(i, name) for i, name in enumerate(NAMES[:20], start=1)]

    # Contoh tanggal bawaan Excel (bukan teks) -> aplikasi mengubahnya ke format Indonesia
    rows[2]["TMT AWAL"] = datetime(2026, 3, 1)

    broken = [_valid_row(i, name) for i, name in enumerate(NAMES[20:], start=21)]
    broken[0]["NO KONTRAK"] = ""                        # 21: No Kontrak kosong
    broken[1]["TMT AKHIR"] = "01 Januari 2026"          # 22: TMT Akhir < TMT Awal
    broken[1]["TMT AWAL"] = "01 Juni 2026"
    broken[2]["UPAH TETAP"] = "lima juta"               # 23: upah bukan angka
    broken[3]["NO KONTRAK"] = rows[0]["NO KONTRAK"]     # 24: No Kontrak dobel dengan data 1
    broken[4]["NIK PEKERJA"] = "999912345"              # 25: NIK tidak 16 digit (peringatan)
    return rows + broken


def build_sample_excel(path: Path) -> None:
    """Menulis file Excel contoh dengan header rapi dan kolom yang lebar."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DATA PKWT"
    sheet.append(HEADERS)
    for row in _sample_rows():
        sheet.append([row[h] for h in HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for column_cells in sheet.columns:
        longest = max(len(str(c.value or "")) for c in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = min(max(12, longest + 2), 45)
    for row_cells in sheet.iter_rows(min_row=2):
        for cell in row_cells:
            if isinstance(cell.value, datetime):
                cell.number_format = "DD/MM/YYYY"
    sheet.freeze_panes = "B2"

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


if __name__ == "__main__":
    build_sample_excel(SAMPLE_EXCEL)
    build_sample_template(SAMPLE_TEMPLATE)
    print(f"OK -> {SAMPLE_EXCEL}")
    print(f"OK -> {SAMPLE_TEMPLATE}")
