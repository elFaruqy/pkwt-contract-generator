"""
Membuat template kontrak PKWT GENERIK (perusahaan fiktif) untuk demo.

Isi pasal sengaja dibuat ringkas dan umum, BUKAN salinan kontrak
perusahaan mana pun. Placeholder {{ ... }} sama persis dengan yang
dipakai aplikasi, jadi template asli kantor tetap kompatibel.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

COMPANY = "PT CONTOH KARYA NUSANTARA"
FONT = "Arial"

# Setiap item: (judul pasal, daftar ayat). Ayat boleh berisi placeholder.
ARTICLES: list[tuple[str, list[str]]] = [
    ("Penempatan dan Jabatan", [
        "PIHAK KEDUA ditempatkan pada pekerjaan {{ NAMA_PROYEK }} (No. IO {{ NO_IO_PROYEK }}) "
        "dengan lokasi kerja di {{ LOKASI_KERJA }}.",
        "PIHAK KEDUA dipekerjakan dengan Nomor Pekerja {{ NOPEK }} dan jabatan {{ JABATAN_PEKERJA }}.",
        "Apabila diperlukan, PIHAK PERTAMA dapat memindahkan PIHAK KEDUA ke bagian atau lokasi "
        "proyek lain dengan pemberitahuan sebelumnya.",
    ]),
    ("Jangka Waktu Perjanjian", [
        "Perjanjian ini berlaku sejak tanggal {{ TMT_AWAL }} sampai dengan tanggal {{ TMT_AKHIR }}.",
        "Perjanjian dapat diperpanjang berdasarkan hasil evaluasi dan kesepakatan PARA PIHAK, "
        "sesuai peraturan ketenagakerjaan yang berlaku.",
    ]),
    ("Upah dan Tunjangan", [
        "PIHAK PERTAMA membayar upah kepada PIHAK KEDUA setiap bulan dengan rincian pada tabel berikut.",
        "Upah dibayarkan paling lambat pada akhir bulan berjalan.",
        "PIHAK KEDUA diikutsertakan dalam program BPJS Ketenagakerjaan dan BPJS Kesehatan "
        "sesuai ketentuan yang berlaku.",
    ]),
    ("Kewajiban Pekerja", [
        "PIHAK KEDUA wajib menaati peraturan perusahaan, peraturan di lokasi proyek, "
        "serta ketentuan keselamatan kerja (HSSE).",
        "PIHAK KEDUA wajib menjaga kerahasiaan informasi perusahaan selama dan setelah masa kerja.",
    ]),
    ("Berakhirnya Perjanjian", [
        "Perjanjian ini berakhir pada tanggal sebagaimana dimaksud dalam Pasal 2, atau lebih awal "
        "karena selesainya pekerjaan proyek atau pelanggaran sesuai peraturan yang berlaku.",
        "Hak dan kewajiban PARA PIHAK saat perjanjian berakhir mengikuti peraturan "
        "perundang-undangan ketenagakerjaan.",
    ]),
    ("Penutup", [
        "Hal-hal yang belum diatur dalam perjanjian ini mengikuti peraturan perusahaan dan "
        "peraturan perundang-undangan yang berlaku.",
        "Perjanjian ini dibuat rangkap 2 (dua) dengan kekuatan hukum yang sama.",
    ]),
]
SALARY_ROWS = [
    ("Upah Tetap", "{{ UPAH_TETAP }}", "Per bulan"),
    ("Tunjangan Daerah", "{{ TUNJANGAN_DAERAH }}", "Per bulan"),
    ("Tunjangan Lainnya", "{{ TUNJANGAN_LAINNYA }}", "Per bulan (sesuai proyek)"),
]
SALARY_ARTICLE_INDEX = 2  # tabel upah disisipkan setelah ayat 1 Pasal 3


def _paragraph(doc, text: str = "", bold: bool = False, align=None, size: int = 11, space_after: int = 6):
    """Menambah paragraf dengan format standar."""
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        para.alignment = align
    if text:
        run = para.add_run(text)
        run.bold = bold
        run.font.size = Pt(size)
    return para


def _numbered(doc, number: int, text: str) -> None:
    """Ayat bernomor dengan indentasi menggantung (hanging indent)."""
    para = _paragraph(doc, f"{number}.\t{text}", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    para.paragraph_format.left_indent = Cm(0.75)
    para.paragraph_format.first_line_indent = Cm(-0.75)
    para.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))


def _set_cell(cell, text: str, bold: bool = False, shade: str | None = None) -> None:
    """Mengisi sel tabel, opsional dengan warna latar."""
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(10.5)
    if shade:
        shading = OxmlElement("w:shd")
        shading.set(qn("w:val"), "clear")
        shading.set(qn("w:fill"), shade)
        cell._tc.get_or_add_tcPr().append(shading)


def _identity_table(doc) -> None:
    """Tabel identitas PIHAK KEDUA."""
    rows = [
        ("Nama", "{{ NAMA_PEKERJA }}"),
        ("Tempat/Tanggal Lahir", "{{ TTL_PEKERJA }}"),
        ("Alamat", "{{ ALAMAT_PEKERJA }}"),
        ("No. KTP", "{{ NIK_PEKERJA }}"),
    ]
    table = doc.add_table(rows=len(rows), cols=3)
    for i, (label, value) in enumerate(rows):
        _set_cell(table.cell(i, 0), label)
        _set_cell(table.cell(i, 1), ":")
        _set_cell(table.cell(i, 2), value, bold=(i == 0))
    table.autofit = False
    for col_index, width in enumerate((Cm(4.5), Cm(0.6), Cm(10.9))):
        table.columns[col_index].width = width
        for cell in table.columns[col_index].cells:
            cell.width = width


def _salary_table(doc) -> None:
    """Tabel rincian upah."""
    table = doc.add_table(rows=1 + len(SALARY_ROWS), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for col, title in enumerate(["No", "Komponen", "Besaran", "Keterangan"]):
        _set_cell(table.cell(0, col), title, bold=True, shade="D9E2F3")
    for i, (component, amount, note) in enumerate(SALARY_ROWS, start=1):
        _set_cell(table.cell(i, 0), str(i))
        _set_cell(table.cell(i, 1), component)
        _set_cell(table.cell(i, 2), amount, bold=True)
        _set_cell(table.cell(i, 3), note)
    _paragraph(doc, space_after=4)


def _signature_table(doc) -> None:
    """Kolom tanda tangan kedua pihak."""
    table = doc.add_table(rows=3, cols=2)
    _set_cell(table.cell(0, 0), "PIHAK PERTAMA", bold=True)
    _set_cell(table.cell(0, 1), "PIHAK KEDUA", bold=True)
    table.rows[1].height = Cm(2.5)
    _set_cell(table.cell(2, 0), "{{ MANAGER }}", bold=True)
    _set_cell(table.cell(2, 1), "{{ NAMA_PEKERJA }}", bold=True)
    for row in table.rows:
        for cell in row.cells:
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _header_footer(doc) -> None:
    """Header penanda dokumen contoh + footer nama pekerja."""
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("DOKUMEN CONTOH — DATA FIKTIF, BUKAN KONTRAK RESMI")
    run.font.size, run.font.color.rgb, run.italic = Pt(8), RGBColor(0x99, 0x99, 0x99), True

    footer = section.footer.paragraphs[0]
    run = footer.add_run("PKWT – {{ NAMA_BAWAH }}")
    run.font.size, run.italic = Pt(8), True


def build_sample_template(path: Path) -> None:
    """Menyusun dan menyimpan template contoh ke `path`."""
    doc = Document()
    section = doc.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Cm(2.5))
    style = doc.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _header_footer(doc)
    _paragraph(doc, "PERJANJIAN KERJA WAKTU TERTENTU", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=14, space_after=2)
    _paragraph(doc, "Nomor: {{ NO_KONTRAK }}", align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    _paragraph(doc, "Pada hari ini, {{ WAKTU_PEMBUATAN_KONTRAK }}, bertempat di Jakarta, "
                    "yang bertanda tangan di bawah ini:", align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _numbered(doc, 1, f"{COMPANY}, berkedudukan di Jakarta, dalam hal ini diwakili oleh {{{{ MANAGER }}}} "
                      "selaku Manager Human Capital, selanjutnya disebut PIHAK PERTAMA.")
    _numbered(doc, 2, "Pekerja dengan identitas berikut, selanjutnya disebut PIHAK KEDUA:")
    _identity_table(doc)
    _paragraph(doc, space_after=2)
    _paragraph(doc, "PARA PIHAK sepakat untuk mengikatkan diri dalam Perjanjian Kerja Waktu Tertentu "
                    "dengan ketentuan sebagai berikut:", align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    for index, (title, clauses) in enumerate(ARTICLES):
        heading = _paragraph(doc, f"PASAL {index + 1}", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.keep_with_next = True
        _paragraph(doc, title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        for number, clause in enumerate(clauses, start=1):
            _numbered(doc, number, clause)
            if index == SALARY_ARTICLE_INDEX and number == 1:
                _salary_table(doc)

    _paragraph(doc, "Demikian perjanjian ini dibuat dan ditandatangani oleh PARA PIHAK dalam keadaan "
                    "sadar dan tanpa paksaan dari pihak mana pun.", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14)
    _signature_table(doc)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
